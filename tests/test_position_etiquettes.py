# -*- coding: utf-8 -*-
"""
Non-régression : chaque étiquette FR d'une planche doit tomber, sur la slide,
à l'emplacement de son mot source dans la page du PDF fabricant — quels que
soient le FORMAT de la page (A3, A2, A4…) et la RÉSOLUTION d'extraction de
l'image sous-jacente.

Bug constaté en prod : `assemble.add_overlay` convertissait les points PDF en
EMU avec une constante (largeur A3, 1190.55 pt) au lieu de la largeur réelle
de la page source. Sur tout autre format, les étiquettes étaient tassées dans
un coin de la planche (page plus petite) ou repoussées hors de la slide (page
plus grande).
"""
from pathlib import Path

import fitz
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from core import assemble, extract, translate

FIXTURE_PDF = Path(__file__).parent / "fixtures" / "DHYA2_test.pdf"
PAGE_PLANCHE = 2  # 1-based : la planche cotée du fixture (page A3 paysage)

pytestmark = pytest.mark.skipif(not FIXTURE_PDF.exists(), reason="fixture PDF absente")


def _pdf_planche_mis_a_l_echelle(tmp_path, echelle: float) -> Path:
    """Copie de la planche du fixture agrandie/réduite d'un facteur `echelle`
    (page ET contenu, en natif : texte toujours extractible et rédigeable)."""
    if echelle == 1:
        return FIXTURE_PDF
    doc = fitz.open(FIXTURE_PDF)
    doc.select([PAGE_PLANCHE - 1])
    page = doc[0]
    page.clean_contents()
    xref = page.get_contents()[0]
    doc.update_stream(xref, f"q {echelle} 0 0 {echelle} 0 0 cm\n".encode() + doc.xref_stream(xref) + b"\nQ")
    rect = fitz.Rect(0, 0, page.rect.width * echelle, page.rect.height * echelle)
    page.set_mediabox(rect)
    page.set_cropbox(rect)
    chemin = tmp_path / f"planche_x{echelle}.pdf"
    doc.save(chemin)
    doc.close()
    return chemin


def _assembler_planche(tmp_path, pdf: Path, page_n: int, dpi: int):
    words = extract.extraire_pdf(pdf, tmp_path, [page_n], dpi=dpi)[page_n]
    planche, _ = translate.construire_planche(page_n, words)
    meta = {"numero": "LDTEST100", "client": "Client", "dessinateur": "AB",
            "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    out = tmp_path / "plan.pptx"
    assemble.assembler({
        "base": assemble.base_pour_type("non accompagné"), "out": out,
        "workdir": tmp_path, "meta": meta, "specs": {}, "planches": [planche],
    })
    return words, planche, Presentation(out)


@pytest.mark.parametrize("dpi", [100, 300])
@pytest.mark.parametrize("echelle", [0.5, 1, 2])
def test_etiquette_a_l_emplacement_de_son_mot_source(tmp_path, echelle, dpi):
    pdf = _pdf_planche_mis_a_l_echelle(tmp_path, echelle)
    page_n = PAGE_PLANCHE if echelle == 1 else 1
    words, planche, prs = _assembler_planche(tmp_path, pdf, page_n, dpi)
    page_w, page_h = words["page_size_pts"]

    slide = prs.slides[2]
    formes = list(slide.shapes)
    i_image = max(i for i, sh in enumerate(formes) if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    image = formes[i_image]
    boites = formes[i_image + 1:]

    labels = planche["labels"]
    assert labels, "le fixture doit produire des étiquettes"
    assert len(boites) == len(labels)

    for lab, tb in zip(labels, boites):
        x0, y0, x1, y1 = lab["bbox"]
        attendu_x = ((x0 + x1) / 2) / page_w
        attendu_y = ((y0 + y1) / 2) / page_h
        reel_x = (tb.left + tb.width / 2 - image.left) / image.width
        reel_y = (tb.top + tb.height / 2 - image.top) / image.height
        if lab.get("ancre") == "debut":
            # Suffixe de cote : l'étiquette est ANCRÉE au début du suffixe (elle
            # ne doit pas mordre sur le nombre qui précède) au lieu d'être
            # centrée. C'est le bord d'ancrage qui doit coïncider, sur l'axe du
            # texte ; l'autre axe reste centré.
            if not lab["vertical"]:
                attendu_x = x0 / page_w
                reel_x = (tb.left - image.left) / image.width
            elif lab["rotation_deg"] == 270:         # texte lu de bas en haut : ancré en bas
                attendu_y = y1 / page_h
                reel_y = (tb.top + tb.height / 2 + tb.width / 2 - image.top) / image.height
            else:                                     # lu de haut en bas : ancré en haut
                attendu_y = y0 / page_h
                reel_y = (tb.top + tb.height / 2 - tb.width / 2 - image.top) / image.height
        assert reel_x == pytest.approx(attendu_x, abs=0.003), (
            f"étiquette « {lab['text']} » (échelle {echelle}, {dpi} dpi) : "
            f"x={reel_x:.3f} au lieu de {attendu_x:.3f}"
        )
        assert reel_y == pytest.approx(attendu_y, abs=0.003), (
            f"étiquette « {lab['text']} » (échelle {echelle}, {dpi} dpi) : "
            f"y={reel_y:.3f} au lieu de {attendu_y:.3f}"
        )


@pytest.mark.parametrize("dpi", [100, 200, 300])
def test_largeur_de_page_de_la_planche_independante_du_dpi(tmp_path, dpi):
    """`page_w_pt` vient de la page source (points PDF), jamais de la
    résolution du rendu : changer le dpi ne doit pas changer l'échelle des
    étiquettes."""
    words = extract.extraire_pdf(FIXTURE_PDF, tmp_path, [PAGE_PLANCHE], dpi=dpi)[PAGE_PLANCHE]
    planche, _ = translate.construire_planche(PAGE_PLANCHE, words)
    with fitz.open(FIXTURE_PDF) as doc:
        assert planche["page_w_pt"] == pytest.approx(doc[PAGE_PLANCHE - 1].rect.width)
