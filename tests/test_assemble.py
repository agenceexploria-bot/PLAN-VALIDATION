# -*- coding: utf-8 -*-
"""
Tests de non-régression pour core/assemble.py — en particulier la purge du
cartouche, dont le point de vigilance n°1 (résidu caché d'un ancien dossier
dans une cellule fusionnée masquée) a été vérifié sur les vrais gabarits
LD82040.pptx / LD64397.pptx (audit de session).
"""
from pptx import Presentation
from PIL import Image

from core import assemble


def _construire(tmp_path, type_equipement="non accompagné", numero="LDTEST100"):
    meta = {"numero": numero, "client": "Nouveau Client SARL", "dessinateur": "AB",
            "indice": "R02", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    out = tmp_path / "test.pptx"
    proj = {
        "base": assemble.base_pour_type(type_equipement), "out": out,
        "workdir": tmp_path, "meta": meta, "specs": {}, "planches": [],
    }
    assemble.assembler(proj)
    return out, meta


def test_purge_cartouche_sur_toutes_les_slides_sans_residu(tmp_path):
    """Aucune cellule du cartouche, sur AUCUNE slide, ne doit conserver un
    résidu d'un ancien dossier (ex. "DIMMER SARL", résidu caché constaté sur
    la base LD82040.pptx dans une cellule fusionnée masquée)."""
    out, meta = _construire(tmp_path)
    prs = Presentation(out)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            for row in shape.table.rows:
                for cell in row.cells:
                    texte = cell.text
                    if "Client" in texte:
                        assert meta["client"] in texte
                        assert "DIMMER" not in texte
                        assert "ECEE" not in texte
                    if texte.strip().upper().startswith("LD") and ":" not in texte:
                        assert texte.strip() == meta["numero"]


def test_deux_contacts_toujours_presents_ld64397(tmp_path):
    """La base LD64397 ('accompagné') n'a qu'un seul contact (Nadia) par
    défaut : ensure_contacts() doit restaurer Jérémie systématiquement."""
    out, _ = _construire(tmp_path, type_equipement="accompagné", numero="LDTEST101")
    prs = Presentation(out)
    texte_garde = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "Nadia" in texte_garde
    assert "Jérémie" in texte_garde


def test_numero_affaire_sur_page_de_garde(tmp_path):
    out, meta = _construire(tmp_path)
    prs = Presentation(out)
    texte_garde = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert meta["numero"] in texte_garde


def test_view3d_sans_planches_ne_crash_pas(tmp_path):
    """Un dossier qui ne fournit que garde + specs + une vue 3D fournisseur
    (aucune planche dessin retenue) ne doit jamais planter. La vue 3D est
    déposée SUR la slide specs (cf. references/cartouche.md) : aucune slide
    supplémentaire n'est créée pour elle — garde + specs = 2 slides."""
    img = tmp_path / "cover_3d_full.png"
    Image.new("RGB", (400, 300), "white").save(img)
    meta = {
        "numero": "LDTEST300", "client": "Client Test", "dessinateur": "AB",
        "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE",
    }
    out = tmp_path / "test_view3d.pptx"
    proj = {
        "base": assemble.base_pour_type("non accompagné"), "out": out,
        "workdir": tmp_path, "meta": meta,
        "specs": {"table": [("Client", "Test")]},
        "view3d": {"image": "cover_3d_full.png", "image_page_w_pt": 800, "callouts": []},
        "planches": [],
    }
    assemble.assembler(proj)  # ne doit lever aucune exception
    prs = Presentation(out)
    assert len(prs.slides) == 2  # garde + specs, aucune slide 3D dédiée


def test_vue3d_deposee_sur_la_slide_specs_avec_planches(tmp_path):
    """La vue 3D fournisseur (avec ses callouts) doit être déposée SUR la
    slide specs elle-même, à côté du tableau FR — jamais sur une slide
    dédiée à part — y compris quand il y a de vraies planches dessin à
    assembler (bug audit : une page 3D dédiée séparée laissait la slide
    specs sans image ni callouts, cf. references/cartouche.md et
    cr-session-2026-06-23.md — décision produit validée avec le client)."""
    img3d = tmp_path / "cover_3d_full.png"
    Image.new("RGB", (400, 300), "white").save(img3d)
    img_planche = tmp_path / "page_2_redacted.png"
    Image.new("RGB", (400, 300), "white").save(img_planche)

    meta = {"numero": "LDTEST600", "client": "Client Test", "dessinateur": "AB",
            "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    out = tmp_path / "test_v3d_specs.pptx"
    proj = {
        "base": assemble.base_pour_type("non accompagné"), "out": out,
        "workdir": tmp_path, "meta": meta,
        "specs": {"table": [("Client", "Test")]},
        "view3d": {"image": "cover_3d_full.png", "image_page_w_pt": 800,
                   "callouts": [{"text": "Structure", "bbox": [0, 0, 20, 10]}]},
        "planches": [{"image": "page_2_redacted.png", "page_n": 2, "labels": []}],
    }
    assemble.assembler(proj)
    prs = Presentation(out)

    assert len(prs.slides) == 3  # garde + specs (3D incluse) + 1 planche, pas de slide en plus

    specs = prs.slides[1]
    images = [sh for sh in specs.shapes if sh.shape_type == 13]
    assert images, "aucune image 3D déposée sur la slide specs"
    tables = [sh for sh in specs.shapes if sh.has_table]
    assert tables, "le tableau FR a disparu de la slide specs"
    textes = " ".join(sh.text_frame.text for sh in specs.shapes if sh.has_text_frame)
    assert "Structure" in textes  # callout FR bien présent sur la slide specs


def test_taille_etiquette_derive_du_texte_et_de_la_bbox():
    """Les étiquettes ne doivent plus avoir une taille fixe identique quel
    que soit le texte (bug audit : 720000x216000 EMU pour toutes les
    étiquettes des planches, la bbox du mot source et la longueur du texte
    FR affiché n'étaient jamais prises en compte) : un texte plus long doit
    produire une boîte plus large, et une bbox source plus haute une boîte
    plus haute."""
    prs = Presentation(assemble.base_pour_type("non accompagné"))
    slide = prs.slides[1]
    # epp réaliste (échelle d'une image déposée sur une slide, disp_w en EMU
    # très supérieur à page_w en points) : sans ça, les planchers minimaux de
    # add_overlay dominent toujours et le test ne distingue rien.
    geom = (0, 0, 7_200_000, 7_200_000)  # image affichée sur 20 cm de large
    page_w = 100  # page source de 100 pt de large
    n_avant = len(slide.shapes)

    assemble.add_overlay(slide, "g", [0, 0, 1, 10], geom, page_w=page_w, fpt=8)
    assemble.add_overlay(slide, "MAGNETOTHERMIC", [0, 0, 1, 10], geom, page_w=page_w, fpt=8)
    assemble.add_overlay(slide, "X", [0, 0, 10, 5], geom, page_w=page_w, fpt=8)
    assemble.add_overlay(slide, "X", [0, 0, 10, 80], geom, page_w=page_w, fpt=8)

    tb_court, tb_long, tb_bas, tb_haut = list(slide.shapes)[n_avant:]
    assert tb_long.width > tb_court.width, "un texte plus long doit produire une boîte plus large"
    assert not (tb_court.width == 720000 and tb_long.width == 720000), (
        "les deux boîtes ne doivent plus retomber sur l'ancienne taille fixe"
    )
    assert tb_haut.height > tb_bas.height, "une bbox source plus haute doit produire une boîte plus haute"


def test_cartouche_cellules_fusionnees_masquees_videes_pas_dupliquees(tmp_path):
    """Les cellules de continuation d'une fusion horizontale/verticale
    (hMerge/vMerge) du cartouche ne doivent jamais dupliquer la valeur de la
    cellule voisine ni conserver un résidu : elles doivent être vides (bug
    audit — constaté sur templates/LD82040.pptx : date de création, client
    et mention légale dupliqués/incohérents entre cellules voisines)."""
    out, meta = _construire(tmp_path, numero="LDTEST700")
    prs = Presentation(out)
    trouve_au_moins_une_cellule_fusionnee = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            for row in shape.table.rows:
                for cell in row.cells:
                    tc = cell._tc
                    if tc.get("hMerge") == "1" or tc.get("vMerge") == "1":
                        trouve_au_moins_une_cellule_fusionnee = True
                        assert cell.text.strip() == "", (
                            f"cellule fusionnée masquée non vidée : {cell.text!r}"
                        )
    assert trouve_au_moins_une_cellule_fusionnee, "fixture de test obsolète : plus aucune cellule fusionnée dans la base"


# ---------------------------------------------------------------------------
# A2 — specs de l'ancien projet retirées de la slide 2, sur les DEUX gabarits
# (audit §4 « Critique » : LD64397 gardait « 1500 kg / 2295 x 2980 / 3340 mm… »
# sous le nouveau tableau, car le retrait ciblait des noms de shapes propres
# à LD82040).
# ---------------------------------------------------------------------------
import copy

import pytest

from core import verify

TYPES = ["non accompagné", "accompagné"]

# Contenu propre aux anciens projets des gabarits (jamais un nom de shape).
RESIDUS_ANCIEN_PROJET = [
    "Fonctionnement", "Capacité de charge", "Machinerie", "Accès paliers",
    "Finition", "Prérequis", "Détails et caractéristiques",
    "1500 kg", "1000 kg", "2295", "2980", "3340", "3930", "2835", "RAL 7016",
]
TITRES_A_CONSERVER = ["PLANS DE VALIDATION", "MONTE-CHARGE"]


def _textes_hors_tableaux(slide):
    """Texte de toutes les shapes texte (groupes inclus) hors tableaux."""
    textes = []

    def parcourir(shapes):
        for sh in shapes:
            if sh.shape_type == 6:  # groupe
                parcourir(sh.shapes)
            elif sh.has_text_frame and sh.text_frame.text.strip():
                textes.append(sh.text_frame.text)
    parcourir(slide.shapes)
    return textes


def _assembler(tmp_path, base, nom="a2.pptx"):
    meta = {"numero": "LDTEST400", "client": "Client A2", "dessinateur": "AB",
            "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    out = tmp_path / nom
    assemble.assembler({
        "base": base, "out": out, "workdir": tmp_path, "meta": meta,
        "specs": {"table": [("Modèle", "NOUVEAU")]}, "planches": [],
    })
    return out


@pytest.mark.parametrize("type_equipement", TYPES)
def test_aucune_spec_d_ancien_projet_sur_la_slide_2(tmp_path, type_equipement):
    out = _assembler(tmp_path, assemble.base_pour_type(type_equipement))
    textes = " | ".join(_textes_hors_tableaux(Presentation(out).slides[1]))
    trouves = [r for r in RESIDUS_ANCIEN_PROJET if r in textes]
    assert not trouves, f"Résidus d'un ancien projet sur la slide 2 ({type_equipement}) : {trouves}"


@pytest.mark.parametrize("type_equipement", TYPES)
def test_titres_de_la_slide_2_conserves(tmp_path, type_equipement):
    out = _assembler(tmp_path, assemble.base_pour_type(type_equipement))
    textes = " | ".join(_textes_hors_tableaux(Presentation(out).slides[1]))
    for titre in TITRES_A_CONSERVER:
        assert titre in textes, f"« {titre} » supprimé à tort ({type_equipement})"


@pytest.mark.parametrize("type_equipement", TYPES)
def test_retrait_des_specs_ne_depend_d_aucun_nom_de_shape(tmp_path, type_equipement):
    """Gabarit dont TOUS les noms de shapes ont été changés : le retrait doit
    se faire sur le contenu, pas sur un nom/id figé."""
    prs = Presentation(assemble.base_pour_type(type_equipement))
    for i, sh in enumerate(prs.slides[1].shapes):
        sh.name = f"Renomme {i}"
    base = tmp_path / "base_renommee.pptx"
    prs.save(base)

    out = _assembler(tmp_path, base)
    textes = " | ".join(_textes_hors_tableaux(Presentation(out).slides[1]))
    assert not [r for r in RESIDUS_ANCIEN_PROJET if r in textes]
    assert "PLANS DE VALIDATION" in textes


@pytest.mark.parametrize("type_equipement", TYPES)
def test_nouveau_tableau_specs_present_sur_la_slide_2(tmp_path, type_equipement):
    out = _assembler(tmp_path, assemble.base_pour_type(type_equipement))
    tables = [sh.table for sh in Presentation(out).slides[1].shapes if sh.has_table]
    cellules = " ".join(c.text for t in tables for r in t.rows for c in r.cells)
    assert "NOUVEAU" in cellules


def test_controle_completude_signale_un_bloc_specs_d_ancien_projet(tmp_path):
    """Filet de sécurité : si un bloc de specs d'ancien projet subsiste (gabarit
    modifié, structure inattendue), le rapport de vérification l'annonce."""
    out = _assembler(tmp_path, assemble.base_pour_type("accompagné"))
    meta = {"numero": "LDTEST400", "client": "Client A2"}
    assert not [a for a in verify.controle_completude(out, meta) if "spécifications" in a]

    prs = Presentation(out)
    tb = prs.slides[1].shapes.add_textbox(0, 0, 3000000, 3000000)
    tb.text_frame.text = "Modèle\nCapacité de charge : 1500 kg\nCourse totale : 3340 mm\nMachinerie"
    prs.save(out)

    alertes = verify.controle_completude(out, meta)
    assert any("Slide 2" in a and "spécifications" in a for a in alertes), alertes
