# -*- coding: utf-8 -*-
"""
Tests de non-régression pour core/verify.py — faux positifs trouvés en
testant l'app de bout en bout (AppTest) sur un vrai dossier.
"""
from pathlib import Path

import pytest

from core import assemble, checklist, verify

FIXTURE_XLSX = Path(__file__).parent / "fixtures" / "checklist_vierge.xlsx"


def test_nombres_deck_exclut_adresse_et_numero_affaire(tmp_path):
    """Bout en bout sur un vrai gabarit : ni l'adresse fixe ni les chiffres
    du numéro d'affaire généré ne doivent apparaître dans `nombres_deck`."""
    meta = {"numero": "LDTEST999", "client": "Client Test", "dessinateur": "AB",
            "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    out = tmp_path / "test.pptx"
    proj = {
        "base": assemble.base_pour_type("non accompagné"), "out": out,
        "workdir": tmp_path, "meta": meta, "specs": {}, "planches": [],
    }
    assemble.assembler(proj)
    deck = verify.nombres_deck(out)
    assert "260" not in deck
    assert "01700" not in deck
    assert "999" not in deck  # chiffres de LDTEST999


def test_adresse_fixe_detectee_pour_exclusion():
    """L'adresse du cartouche (fixe, jamais une cote fabricant) doit être
    reconnue par `CARTOUCHE_FIXE` : c'est ce qui permet à `nombres_deck` de
    l'exclure entièrement du comptage. Bug trouvé : '260' et '01700'
    remontaient comme valeurs 'inventées' (BLOQUANT) à chaque génération,
    car ce filtre n'existait pas encore."""
    texte = "260 rue des Barronnières\n01700 - BEYNOST\nTel. +33 4 78 88 14 00"
    assert verify.CARTOUCHE_FIXE.search(texte)


def test_numero_affaire_exclu_du_comptage():
    """Les chiffres du numéro d'affaire généré (LDxxxxx) ne sont jamais une
    cote fabricant. Bug trouvé : 'LDTEST004' faisait remonter '004' comme
    valeur inventée (BLOQUANT)."""
    assert verify.numbers_in("LDTEST004") == {}
    # mais un nombre à côté du numéro d'affaire reste bien détecté
    assert verify.numbers_in("LDTEST004 - 1000 kg") == {"1000": 1}


def test_libelles_cartouche_exclus():
    for texte in ["Dessinateur : AB", "Date de création : 17/09/2026",
                  "Plan de validation indice : R00", "Client : Test",
                  "Page : 3"]:
        assert verify.CARTOUCHE_FIXE.search(texte), f"non reconnu : {texte!r}"


@pytest.mark.skipif(not FIXTURE_XLSX.exists(), reason="fixture checklist absente")
def test_checklist_vierge_reelle_non_utilisee():
    """La checklist réelle fournie par Vertical est le modèle vierge — elle
    ne doit jamais servir de source de comparaison (cf. SKILL.md)."""
    lignes = checklist.lire_checklist(FIXTURE_XLSX)
    assert checklist.est_vierge(lignes) is True


def test_checklist_invalide_ne_crash_pas(tmp_path):
    """Un fichier .xlsx sans feuille 'Technique' (donc pas la checklist
    attendue) ne doit jamais faire planter controle_cotes() (bug audit :
    ValueError non rattrapée) — il doit être ignoré comme source de
    comparaison, avec un message clair pour l'utilisateur."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Feuille1"
    chemin = tmp_path / "checklist_invalide.xlsx"
    wb.save(chemin)

    pptx_path = assemble.base_pour_type("non accompagné")
    rapport = verify.controle_cotes({}, pptx_path, chemin)  # ne doit lever aucune exception
    assert rapport["checklist"]["invalide"] is True
    assert "Technique" in rapport["checklist"]["message"]


def test_bloc_specs_residuel_detecte_par_completude(tmp_path):
    """Filet de sécurité de `controle_completude` (indépendant du retrait fait
    à l'assemblage par `assemble.retirer_blocs_specs`) : si un gabarit change
    de structure et qu'un bloc de specs de l'ancien projet subsiste malgré
    tout sur une slide, le contrôle de complétude doit le voir plutôt que de
    laisser un « 1500 kg » d'un autre dossier passer inaperçu."""
    meta = {"numero": "LDTEST300", "client": "Client Test", "dessinateur": "AB",
            "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    out = tmp_path / "test.pptx"
    proj = {
        "base": assemble.base_pour_type("non accompagné"), "out": out,
        "workdir": tmp_path, "meta": meta, "specs": {"table": [("Modèle", "TEST")]}, "planches": [],
    }
    assemble.assembler(proj)

    # Cas normal : rien à signaler (le retrait à l'assemblage a fait son travail).
    assert verify.controle_completude(out, meta) == []

    # On réinjecte, à la main, un bloc de specs « ancien projet » (simule un
    # gabarit qui n'aurait pas été nettoyé) : deux intertitres distincts.
    from pptx import Presentation
    from pptx.util import Cm

    prs = Presentation(out)
    box = prs.slides[0].shapes.add_textbox(Cm(1), Cm(1), Cm(5), Cm(3))
    box.text_frame.text = "Modèle : ANCIEN\nFinition : RAL 9001"
    prs.save(out)

    alertes = verify.controle_completude(out, meta)
    assert any("bloc de spécifications d'un ancien projet" in a for a in alertes)


def test_page_specs_vide_signalee_par_completude(tmp_path):
    """Constaté au premier test réel sur Render : page 2 (specs) sans vue 3D
    ni tableau FR, passée sans aucune alerte. La page specs doit toujours
    porter au moins le tableau FR reconstruit."""
    meta = {"numero": "LDTEST301", "client": "Client Test", "dessinateur": "AB",
            "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    out = tmp_path / "vide.pptx"
    assemble.assembler({
        "base": assemble.base_pour_type("non accompagné"), "out": out,
        "workdir": tmp_path, "meta": meta, "specs": {}, "planches": [],
    })
    alertes = verify.controle_completude(out, meta)
    assert any("page specs" in a.lower() and "vide" in a.lower() for a in alertes), alertes


def _checklist_non_vierge(tmp_path, lignes):
    """Construit une checklist .xlsx réelle, feuille « Technique », avec des
    valeurs renseignées (par opposition au modèle vierge — cf.
    `checklist.est_vierge`)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Technique"
    ws.append(["Libellé", "Réponse"])
    for libelle, valeur in lignes:
        ws.append([libelle, valeur])
    chemin = tmp_path / "checklist_remplie.xlsx"
    wb.save(chemin)
    return chemin


def test_checklist_non_vierge_sert_de_source_de_comparaison(tmp_path):
    """Quand la checklist est réellement renseignée (pas le modèle vierge),
    `controle_cotes` doit s'en servir : une valeur de la checklist absente à
    la fois de la source PDF et du deck généré doit remonter comme écart à
    arbitrer ; une valeur présente dans la source ne doit PAS y remonter
    (faux positif)."""
    checklist_path = _checklist_non_vierge(
        tmp_path, [("Capacité", 1500), ("Vitesse", "0.15")],
    )
    meta = {"numero": "LDTEST400", "client": "Client Test", "dessinateur": "AB",
            "indice": "R00", "date": "01/01/2026", "equipement": "MONTE-CHARGE"}
    pptx_path = tmp_path / "test.pptx"
    assemble.assembler({
        "base": assemble.base_pour_type("non accompagné"), "out": pptx_path,
        "workdir": tmp_path, "meta": meta, "specs": {}, "planches": [],
    })
    words_par_page = {1: {"words": [{"text": "Vitesse 0.15 m/s"}]}}

    rapport = verify.controle_cotes(words_par_page, pptx_path, checklist_path)

    cl = rapport["checklist"]
    assert cl["vierge"] is False
    assert "1500" in cl["valeurs_checklist_absentes_du_dossier"]
    assert "0.15" not in cl["valeurs_checklist_absentes_du_dossier"]
    assert cl["lignes_technique"]["Capacité"]["Réponse"] == 1500


def test_pages_scannees_empeche_le_pass_trompeur(tmp_path):
    """Un contrôle de cotes réalisé alors que des pages retenues sont sans
    couche texte (PDF scanné) ne doit jamais afficher PASS, même si les
    compteurs source/deck sont égaux par coïncidence (bug audit : 0 source
    vs 0 deck remontait PASS à tort) : statut distinct dédié."""
    pptx_path = assemble.base_pour_type("non accompagné")
    rapport = verify.controle_cotes({}, pptx_path, pages_scannees=[1])
    assert rapport["statut_cotes"] == "NON VÉRIFIABLE (PDF scanné)"
    assert rapport["pages_scannees"] == [1]
