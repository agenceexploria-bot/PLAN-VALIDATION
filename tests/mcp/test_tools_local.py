# -*- coding: utf-8 -*-
"""
Test local des 6 outils MCP, appelés directement comme des fonctions Python
(pas de client/serveur MCP réel ici — cf. tests/mcp/test_server_http.py pour
un échange MCP complet, protocole réel). Pipeline complet sur
tests/fixtures/DHYA2_test.pdf : inventaire -> extraction (garde + planche)
-> traduction -> assemblage -> vérification -> export PDF (refus sans
validation, puis avec).

Les outils ne renvoient jamais de contenu binaire inline (base64) mais une
URL de téléchargement à usage unique (cf. mcp_server/fichiers.py — le SDK
MCP plafonne une réponse d'outil à 1 Mio, dépassé par un vrai PPTX/planche).
`_recuperer_base64` simule ce que ferait un vrai client : un GET sur l'URL,
puis un ré-encodage en base64 pour le passer à l'outil suivant. Ici, appel
direct au registre plutôt qu'un GET HTTP réel (pas de serveur dans ces
tests) — le téléchargement HTTP réel est couvert par test_server_http.py.
"""
import base64
from pathlib import Path

import pytest

from mcp_server import fichiers
from mcp_server.tools import assemblage, export, extraction, inventaire, traduction, verification

FIXTURE_PDF = Path(__file__).resolve().parent.parent / "fixtures" / "DHYA2_test.pdf"
pytestmark = pytest.mark.skipif(not FIXTURE_PDF.exists(), reason="fixture PDF absente")


def _jeton_de(url: str) -> str:
    return url.rsplit("/", 1)[-1]


def _contenu_publie(url: str) -> bytes:
    resultat = fichiers.recuperer_et_invalider(_jeton_de(url))
    assert resultat is not None, f"fichier publié introuvable : {url}"
    return resultat["contenu"]


def _recuperer_base64(url: str) -> str:
    return base64.b64encode(_contenu_publie(url)).decode("ascii")


@pytest.fixture(scope="module")
def pdf_base64():
    return base64.b64encode(FIXTURE_PDF.read_bytes()).decode("ascii")


@pytest.fixture(scope="module")
def pdf_id(pdf_base64):
    """`inventaire_pdf` n'est appelé qu'UNE fois par pipeline (cf.
    mcp_server/pdf_cache.py) — les tests suivants réutilisent le pdf_id
    qu'il retourne, comme le ferait l'agent Dust réel, au lieu de
    retransmettre le PDF en base64 à chaque outil."""
    return inventaire.inventaire_pdf(pdf_base64=pdf_base64)["pdf_id"]


def test_inventaire_pdf(pdf_base64):
    resultat = inventaire.inventaire_pdf(pdf_base64=pdf_base64)
    assert resultat["pdf_id"]
    assert resultat["n_pages"] == 2
    assert len(resultat["pages"]) == 2
    for page in resultat["pages"]:
        assert page["apercu_url"].startswith("http")
        assert _contenu_publie(page["apercu_url"])[:8] == b"\x89PNG\r\n\x1a\n"
        assert len(page["page_size_pts"]) == 2


def test_inventaire_pdf_exige_exactement_un_des_deux(pdf_base64, pdf_id):
    with pytest.raises(ValueError, match="pdf_base64 ou pdf_id"):
        inventaire.inventaire_pdf()
    with pytest.raises(ValueError, match="pdf_base64 ou pdf_id"):
        inventaire.inventaire_pdf(pdf_base64=pdf_base64, pdf_id=pdf_id)


def test_extraire_page_garde(pdf_id):
    resultat = extraction.extraire_page(pdf_id=pdf_id, page_num=1, dpi=150, role="garde")
    assert resultat["page_num"] == 1
    assert resultat["pdf_id"] == pdf_id
    assert "image_url" in resultat
    assert "survie_nombres" not in resultat  # role=garde : rediger=False, aucune rédaction
    # La vue 3D peut échouer à s'extraire sur ce fixture minimal (pas de
    # tableau specs distinct à masquer) — les deux issues sont acceptables,
    # seul un plantage de l'outil ne le serait pas.
    assert "vue_3d_url" in resultat or "vue_3d_erreur" in resultat


def test_extraire_page_planche(pdf_id):
    resultat = extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=150, role="planche")
    assert resultat["page_num"] == 2
    assert _contenu_publie(resultat["image_url"])[:8] == b"\x89PNG\r\n\x1a\n"
    assert resultat["survie_nombres"]["perdus"] == []
    assert resultat["redaction_pil_fallback"] == []
    assert any(w["translatable"] for w in resultat["words"])


def test_extraire_page_planche_via_base64_direct(pdf_base64):
    """Repli base64 (petits fichiers / tests directs) — cf.
    util.resoudre_pdf. Le chemin normal pour un agent Dust réel reste
    pdf_id, testé par ailleurs."""
    resultat = extraction.extraire_page(pdf_base64=pdf_base64, page_num=2, dpi=150, role="planche")
    assert resultat["page_num"] == 2
    assert resultat["pdf_id"]  # un nouveau pdf_id a été mis en cache


def test_extraire_page_role_invalide(pdf_id):
    with pytest.raises(ValueError):
        extraction.extraire_page(pdf_id=pdf_id, page_num=1, role="inconnu")


def test_extraire_page_hors_limites(pdf_id):
    with pytest.raises(ValueError):
        extraction.extraire_page(pdf_id=pdf_id, page_num=99)


def test_extraire_page_pdf_id_inconnu_refuse():
    with pytest.raises(ValueError, match="pdf_id"):
        extraction.extraire_page(pdf_id="jeton-inexistant", page_num=1)


def test_pdf_trop_volumineux_refuse():
    faux_pdf_base64 = base64.b64encode(b"x" * (51 * 1_000_000)).decode("ascii")
    with pytest.raises(ValueError, match="volumineux"):
        inventaire.inventaire_pdf(pdf_base64=faux_pdf_base64)


def test_traduire_mots_planche(pdf_id):
    """Chemin NORMAL pour un agent Dust réel : extraction_id, pas de JSON
    words_data reconstruit à la main (cf. mcp_server/extraction_cache.py)."""
    page = extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=150, role="planche")
    assert page["extraction_id"]
    resultat = traduction.traduire_mots(extraction_id=page["extraction_id"], role="planche")
    assert isinstance(resultat["labels"], list) and resultat["labels"]
    assert resultat["page_w_pt"] == page["page_size_pts"][0]


def test_traduire_mots_planche_via_words_data_direct(pdf_id):
    """Repli words_data (petits fichiers / tests directs) — cf.
    util.resoudre_words_data. Le chemin normal pour un agent Dust réel reste
    extraction_id, testé par ailleurs."""
    page = extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=150, role="planche")
    resultat = traduction.traduire_mots(
        words_data={"page_size_pts": page["page_size_pts"], "words": page["words"]}, role="planche",
    )
    assert isinstance(resultat["labels"], list) and resultat["labels"]


def test_traduire_mots_exige_exactement_un_des_deux(pdf_id):
    page = extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=150, role="planche")
    with pytest.raises(ValueError, match="words_data ou extraction_id"):
        traduction.traduire_mots(role="planche")
    with pytest.raises(ValueError, match="words_data ou extraction_id"):
        traduction.traduire_mots(
            extraction_id=page["extraction_id"],
            words_data={"page_size_pts": page["page_size_pts"], "words": page["words"]},
            role="planche",
        )


def test_traduire_mots_extraction_id_inconnu_refuse():
    with pytest.raises(ValueError, match="extraction_id"):
        traduction.traduire_mots(extraction_id="jeton-inexistant", role="planche")


def test_traduire_mots_words_data_malformee_message_clair(pdf_id):
    """Bug corrigé : un words_data mal formé (ex. reconstruit à la main par
    un agent qui a buté sur la taille du JSON complet) doit renvoyer une
    erreur explicite nommant le champ en cause — jamais un KeyError/
    TypeError opaque, jamais un échec silencieux."""
    with pytest.raises(ValueError, match="page_size_pts"):
        traduction.traduire_mots(words_data={"words": []}, role="planche")

    with pytest.raises(ValueError, match="'words' manquant"):
        traduction.traduire_mots(words_data={"page_size_pts": [100, 200]}, role="planche")

    with pytest.raises(ValueError, match=r"words\[0\]"):
        traduction.traduire_mots(
            words_data={"page_size_pts": [100, 200], "words": [{"text": "X"}]}, role="planche",
        )


def test_traduire_mots_glossaire_version_non_supportee(pdf_id):
    page = extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=150, role="planche")
    with pytest.raises(ValueError, match="glossaire_version"):
        traduction.traduire_mots(
            extraction_id=page["extraction_id"], role="planche", glossaire_version="2026-01",
        )


# ---------------------------------------------------------------------------
# Pipeline complet : assemblage -> vérification -> export
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def projet_assemble(pdf_id):
    page_planche = extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=150, role="planche")
    image_planche_b64 = _recuperer_base64(page_planche["image_url"])
    trad_planche = traduction.traduire_mots(extraction_id=page_planche["extraction_id"], role="planche")

    projet_json = {
        "meta": {
            "numero": "LDMCP001", "client": "Client Test MCP", "dessinateur": "AB",
            "indice": "R00", "type_equipement": "non accompagné",
        },
        "specs": {"table": [["Modèle", "TEST"]]},
        "planches": [{
            "image_base64": image_planche_b64,
            "page_n": 2,
            "page_w_pt": page_planche["page_size_pts"][0],
            "labels": trad_planche["labels"],
        }],
        "hors_glossaire": trad_planche["hors_glossaire"],
    }
    resultat = assemblage.assembler_pptx(projet_json)
    pptx_base64 = _recuperer_base64(resultat["pptx_url"])  # jeton consommé une seule fois, ici
    return resultat, pptx_base64, {2: page_planche}


def test_assembler_pptx(projet_assemble):
    resultat, pptx_base64, _ = projet_assemble
    assert resultat["pptx_sha256"]
    assert base64.b64decode(pptx_base64)[:2] == b"PK"  # signature ZIP/OOXML
    assert resultat["resume"]["nb_slides"] >= 3  # garde + specs + 1 planche
    assert resultat["resume"]["nb_planches"] == 1


def test_assembler_pptx_meta_incomplete_refuse():
    with pytest.raises(ValueError, match="numero"):
        assemblage.assembler_pptx({"meta": {"client": "X", "dessinateur": "AB",
                                             "indice": "R00", "type_equipement": "non accompagné"}})


def test_assembler_pptx_numero_mal_forme_refuse():
    with pytest.raises(ValueError, match="LDxxxxx"):
        assemblage.assembler_pptx({"meta": {
            "numero": "PASBONDUTOUT", "client": "X", "dessinateur": "AB",
            "indice": "R00", "type_equipement": "non accompagné",
        }})


META_VALIDE = {"numero": "LDMCP004", "client": "X", "dessinateur": "AB",
               "indice": "R00", "type_equipement": "non accompagné"}


@pytest.fixture(scope="module")
def planche_300dpi(pdf_id):
    """Planche A3 à 300 dpi (4961×3508 px) : la taille RÉELLE du cas qui a
    fait échouer un agent Dust (image en base64 ≈ 165k jetons)."""
    return extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=300, role="planche")


def test_assembler_pptx_par_extraction_id_seulement(planche_300dpi):
    """Chemin NORMAL : la planche passe par son seul extraction_id — le
    serveur retrouve l'image pleine résolution, la largeur de page et les
    labels de traduire_mots. Le pptx_id retourné résout le PPTX produit."""
    import io
    import zipfile
    from PIL import Image
    from mcp_server import util

    eid = planche_300dpi["extraction_id"]
    trad = traduction.traduire_mots(extraction_id=eid, role="planche")
    resultat = assemblage.assembler_pptx({
        "meta": META_VALIDE, "specs": {"table": [["Modèle", "TEST"]]},
        "planches": [{"extraction_id": eid}],
    })
    pptx = util.resoudre_pptx(None, resultat["pptx_id"])
    assert pptx == _contenu_publie(resultat["pptx_url"])
    with zipfile.ZipFile(io.BytesIO(pptx)) as z:
        tailles = [Image.open(io.BytesIO(z.read(n))).size
                   for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".png")]
        xml_planche = z.read("ppt/slides/slide3.xml").decode("utf-8")
    assert (4961, 3508) in tailles, tailles
    assert trad["labels"][0]["text"] in xml_planche  # labels repris sans retransmission


def test_assembler_pptx_sans_specs_refuse(planche_300dpi):
    """Premier test réel sur Render : page specs sortie VIDE (ni vue 3D ni
    tableau), sans aucune erreur. Le tableau specs est obligatoire."""
    eid = planche_300dpi["extraction_id"]
    traduction.traduire_mots(extraction_id=eid, role="planche")
    for specs in (None, {}, {"table": []}):
        projet = {"meta": META_VALIDE, "planches": [{"extraction_id": eid}]}
        if specs is not None:
            projet["specs"] = specs
        with pytest.raises(ValueError, match=r"specs.*traduire_mots\(extraction_id=\.\.\., role='specs'\)"):
            assemblage.assembler_pptx(projet)


def test_assembler_pptx_specs_par_extraction_id(pdf_id, planche_300dpi):
    """Comme les labels : le tableau specs produit par traduire_mots(role='specs')
    est rattaché à l'extraction, l'agent n'a pas à le retransmettre."""
    from mcp_server import util
    from pptx import Presentation
    import io

    page_specs = extraction.extraire_page(pdf_id=pdf_id, page_num=1, dpi=150, role="specs")
    table = traduction.traduire_mots(extraction_id=page_specs["extraction_id"], role="specs")["table"]
    eid = planche_300dpi["extraction_id"]
    traduction.traduire_mots(extraction_id=eid, role="planche")
    resultat = assemblage.assembler_pptx({
        "meta": META_VALIDE, "specs": {"extraction_id": page_specs["extraction_id"]},
        "planches": [{"extraction_id": eid}],
    })
    prs = Presentation(io.BytesIO(util.resoudre_pptx(None, resultat["pptx_id"])))
    textes = [c.text for sh in prs.slides[1].shapes if sh.has_table for r in sh.table.rows for c in r.cells]
    assert table[0][0] in textes


def test_assembler_pptx_vue_3d_et_specs_meme_page(pdf_id, planche_300dpi, capsys):
    """2e test Dust réel : page specs sans vue 3D. Scénario contrôlé du
    chemin documenté : page 1 (vue 3D + tableau specs) extraite UNE fois avec
    role='garde', traduite avec role='specs', même extraction_id pour view3d
    et specs -> la vue 3D est bien déposée sur la page specs. Le log serveur
    trace la présence de chaque entrée, sans identifiant."""
    from mcp_server import util
    from pptx import Presentation
    import io

    page1 = extraction.extraire_page(pdf_id=pdf_id, page_num=1, dpi=150, role="garde")
    eid1 = page1["extraction_id"]
    traduction.traduire_mots(extraction_id=eid1, role="specs")
    eid = planche_300dpi["extraction_id"]
    traduction.traduire_mots(extraction_id=eid, role="planche")
    capsys.readouterr()
    resultat = assemblage.assembler_pptx({
        "meta": META_VALIDE, "view3d": {"extraction_id": eid1}, "specs": {"extraction_id": eid1},
        "planches": [{"extraction_id": eid}],
    })
    log = capsys.readouterr().out
    assert "[assembler_pptx] view3d=extraction_id specs=extraction_id(vue_3d=oui) planches=1 [extraction_id]" in log
    assert eid1 not in log and eid not in log
    prs = Presentation(io.BytesIO(util.resoudre_pptx(None, resultat["pptx_id"])))
    assert any(sh.shape_type == 13 for sh in prs.slides[1].shapes)  # PICTURE


@pytest.mark.parametrize("role", ["garde", "specs"])
def test_assembler_pptx_vue_3d_deduite_de_specs(pdf_id, planche_300dpi, role, capsys):
    """3e test Dust réel : view3d=ABSENT malgré les instructions. Page 1
    (vue 3D + tableau specs) extraite UNE seule fois, quel que soit le rôle
    choisi, et passée seulement en specs : la vue 3D en est déduite, posée
    sur la page specs, et SIGNALÉE à l'agent pour qu'il prévienne
    l'utilisateur avant verifier_rendu."""
    from mcp_server import util
    from pptx import Presentation
    import io

    page1 = extraction.extraire_page(pdf_id=pdf_id, page_num=1, dpi=150, role=role)
    assert "vue_3d_url" in page1
    eid1 = page1["extraction_id"]
    traduction.traduire_mots(extraction_id=eid1, role="specs")
    eid = planche_300dpi["extraction_id"]
    traduction.traduire_mots(extraction_id=eid, role="planche")
    capsys.readouterr()
    resultat = assemblage.assembler_pptx({
        "meta": META_VALIDE, "specs": {"extraction_id": eid1}, "planches": [{"extraction_id": eid}],
    })
    assert resultat["resume"]["vue_3d_deduite"] is True
    assert any("déduite automatiquement" in m and "vérifier visuellement" in m
               for m in resultat["a_signaler_a_l_utilisateur"])
    assert "specs=extraction_id(vue_3d=oui)" in capsys.readouterr().out
    prs = Presentation(io.BytesIO(util.resoudre_pptx(None, resultat["pptx_id"])))
    assert any(sh.shape_type == 13 for sh in prs.slides[1].shapes)  # PICTURE


def test_assembler_pptx_view3d_explicite_prioritaire(pdf_id, planche_300dpi):
    """view3d fourni explicitement : rien n'est déduit, aucun avertissement."""
    garde = extraction.extraire_page(pdf_id=pdf_id, page_num=1, dpi=150, role="garde")
    specs = extraction.extraire_page(pdf_id=pdf_id, page_num=1, dpi=150, role="specs")
    traduction.traduire_mots(extraction_id=specs["extraction_id"], role="specs")
    eid = planche_300dpi["extraction_id"]
    traduction.traduire_mots(extraction_id=eid, role="planche")
    resultat = assemblage.assembler_pptx({
        "meta": META_VALIDE, "view3d": {"extraction_id": garde["extraction_id"]},
        "specs": {"extraction_id": specs["extraction_id"]}, "planches": [{"extraction_id": eid}],
    })
    assert resultat["resume"]["vue_3d_deduite"] is False
    assert resultat["a_signaler_a_l_utilisateur"] == []


def test_assembler_pptx_journalise_view3d_absent(capsys):
    """Appel refusé (specs manquant) : la présence/absence des entrées est
    tracée quand même, avant la validation."""
    with pytest.raises(ValueError):
        assemblage.assembler_pptx({"meta": META_VALIDE, "planches": ["pas un objet"]})
    assert "view3d=ABSENT specs=ABSENT planches=1 [invalide(str)]" in capsys.readouterr().out


def test_assembler_pptx_sans_traduire_mots_refuse(pdf_id):
    page = extraction.extraire_page(pdf_id=pdf_id, page_num=2, dpi=150, role="planche")
    with pytest.raises(ValueError, match=r"planches\[0\].*traduire_mots"):
        assemblage.assembler_pptx({"meta": META_VALIDE, "planches": [{"extraction_id": page["extraction_id"]}]})


def test_assembler_pptx_extraction_garde_comme_planche_refuse(pdf_id):
    garde = extraction.extraire_page(pdf_id=pdf_id, page_num=1, dpi=150, role="garde")
    with pytest.raises(ValueError, match=r"planches\[0\]\.extraction_id.*role='planche'"):
        assemblage.assembler_pptx({"meta": META_VALIDE, "planches": [{"extraction_id": garde["extraction_id"]}]})


@pytest.mark.parametrize("planche, attendu", [
    ({"page_n": 2}, r"planches\[0\] : fournir exactement un"),
    ({"extraction_id": "jeton-inexistant"}, r"planches\[0\]\.extraction_id .* inconnu ou expiré"),
    ({"image_base64": "pas du base64 !"}, r"planches\[0\]\.image_base64 : base64 invalide"),
    ({"image_base64": "iVBORw0KGgo=", "labels": [{"text": "x", "bbox": [1, 2]}]}, r"planches\[0\]\.labels\[0\]\.bbox"),
    ({"image_base64": "iVBORw0KGgo=", "page_w_pt": "A3"}, r"planches\[0\]\.page_w_pt"),
])
def test_assembler_pptx_entree_malformee_nomme_le_champ(planche, attendu):
    """Bug d'origine : assembler_pptx échouait sur un KeyError/binascii.Error
    opaque. Chaque entrée malformée doit nommer le champ en cause."""
    with pytest.raises(ValueError, match=attendu):
        assemblage.assembler_pptx({"meta": META_VALIDE, "planches": [planche]})


def test_verifier_rendu_et_exporter_pdf_pptx_id_inconnu_refuse():
    with pytest.raises(ValueError, match="pptx_id .* relancez assembler_pptx"):
        verification.verifier_rendu(pptx_id="jeton-inexistant")
    with pytest.raises(ValueError, match="pptx_id .* relancez assembler_pptx"):
        export.exporter_pdf(pptx_id="jeton-inexistant", valide=True)
    with pytest.raises(ValueError, match="exactement un de pptx_id ou pptx_base64"):
        verification.verifier_rendu()


@pytest.mark.parametrize("exception", [
    RuntimeError("LibreOffice : Échec de la conversion (code 1)"),
    OSError("LibreOffice : soffice introuvable"),  # autre type que RuntimeError (ex. com_error)
])
def test_echec_du_moteur_de_rendu_lisible_par_l_agent(monkeypatch, exception):
    """Même bug que les ValueError : un échec du moteur de rendu (seul
    LibreOffice sous Docker/Render) arrivait à l'agent réduit à « Error
    executing tool ... ». Il doit lever une ToolError — seul type dont le SDK
    mcp transmet le texte — portant le détail du moteur."""
    from mcp.server.mcpserver.exceptions import ToolError
    from core import render

    def _echoue(*args, **kwargs):
        raise exception

    monkeypatch.setattr(render, "rendre_pngs", _echoue)
    monkeypatch.setattr(render, "exporter_pdf", _echoue)
    pptx_b64 = base64.b64encode(b"PK\x03\x04 pptx factice").decode("ascii")

    with pytest.raises(ToolError, match=r"verifier_rendu : échec du moteur de rendu.*LibreOffice"):
        verification.verifier_rendu(pptx_b64)
    with pytest.raises(ToolError, match=r"exporter_pdf : échec du moteur de rendu.*LibreOffice"):
        export.exporter_pdf(pptx_b64, valide=True, moteur="libreoffice")


@pytest.fixture(scope="module")
def rendu_verifie(projet_assemble):
    """words_par_page passe l'extraction_id de chaque page, PAS le words_data
    complet — c'est le pire cas du problème corrigé (agréger PLUSIEURS pages
    en JSON aurait été encore plus volumineux que le cas à une page de
    traduire_mots), cf. mcp_server/extraction_cache.py."""
    _, pptx_base64, words_par_page = projet_assemble
    words_par_page_json = {str(n): wd["extraction_id"] for n, wd in words_par_page.items()}
    resultat = verification.verifier_rendu(
        pptx_base64, words_par_page=words_par_page_json,
        meta={"numero": "LDMCP001", "client": "Client Test MCP"},
    )
    return pptx_base64, resultat


def test_verifier_rendu(rendu_verifie):
    _, resultat = rendu_verifie
    assert resultat["moteur"] in ("powerpoint", "libreoffice")
    assert len(resultat["slides"]) >= 3
    for slide in resultat["slides"]:
        assert _contenu_publie(slide["url"])[:8] == b"\x89PNG\r\n\x1a\n"
    assert resultat["pptx_sha256"]
    assert resultat["rapport_cotes"] is not None
    assert resultat["rapport_cotes"]["survie_nombres"]["statut"] == "PASS"


def test_verifier_rendu_words_par_page_extraction_id_inconnu_refuse(projet_assemble):
    _, pptx_base64, _ = projet_assemble
    with pytest.raises(ValueError, match="extraction_id"):
        verification.verifier_rendu(pptx_base64, words_par_page={"2": "jeton-inexistant"})


def test_exporter_pdf_refuse_sans_validation(rendu_verifie):
    pptx_base64, _ = rendu_verifie
    with pytest.raises(ValueError, match="[Rr]efus"):
        export.exporter_pdf(pptx_base64, valide=False)


def test_exporter_pdf_refuse_si_empreinte_ne_correspond_pas(rendu_verifie):
    pptx_base64, _ = rendu_verifie
    with pytest.raises(ValueError, match="empreinte"):
        export.exporter_pdf(pptx_base64, valide=True, pptx_sha256_verifie="0" * 64)


def test_exporter_pdf_valide_avec_empreinte_correcte(rendu_verifie):
    pptx_base64, resultat_verif = rendu_verifie
    resultat = export.exporter_pdf(
        pptx_base64, valide=True,
        pptx_sha256_verifie=resultat_verif["pptx_sha256"],
        moteur=resultat_verif["moteur"],
    )
    assert resultat["pdf_sha256"]
    assert resultat["nom_fichier"].endswith(".pdf")
    assert _contenu_publie(resultat["pdf_url"])[:5] == b"%PDF-"
