# -*- coding: utf-8 -*-
"""
Cartouche rendu par LibreOffice (cf. core/cartouche_libreoffice.py) : sur les
gabarits intacts, le cartouche débordait de 48 à 64 pt sous la page et la
mention légale était absente du PDF. Les tests portent sur le PDF réellement
produit à partir de la copie jetable (`util.copie_allegee`), pour les deux
gabarits.
"""
import hashlib
import zipfile
from pathlib import Path

import fitz
import pytest
from lxml import etree

from core import cartouche_libreoffice as cl
from core import render
from core.assemble import BASES_PAR_TYPE, bornes_verticales
from mcp_server import util

GABARITS = sorted(BASES_PAR_TYPE.values())
A, P = cl.A, cl.P


def _copie(gabarit: Path, dossier: Path) -> Path:
    copie = util.copie_allegee(gabarit, dossier / gabarit.name, 2600)
    assert copie != gabarit
    return copie


@pytest.fixture(scope="module", params=GABARITS, ids=lambda p: p.stem)
def pdf_du_gabarit(request, tmp_path_factory):
    dossier = tmp_path_factory.mktemp(request.param.stem)
    avant = hashlib.sha256(request.param.read_bytes()).hexdigest()
    copie = _copie(request.param, dossier)
    pdf = render.exporter_pdf(copie, moteur=render.MOTEUR_LIBREOFFICE)["pdf"]
    assert hashlib.sha256(request.param.read_bytes()).hexdigest() == avant, \
        "le gabarit ne doit jamais être modifié"
    return request.param, copie, pdf


def _pages_a_cartouche(pdf):
    doc = fitz.open(pdf)
    try:
        # La garde (page 1) n'a pas de cartouche.
        for page in list(doc)[1:]:
            yield page
    finally:
        doc.close()


def test_aucun_trait_ni_texte_hors_page(pdf_du_gabarit):
    _, _, pdf = pdf_du_gabarit
    for page in _pages_a_cartouche(pdf):
        hauteur = page.rect.height
        bas_traits = max(d["rect"].y1 for d in page.get_drawings())
        bas_texte = max(l["bbox"][3] for b in page.get_text("dict")["blocks"] for l in b.get("lines", []))
        assert bas_traits <= hauteur, f"page {page.number + 1} : trait à {bas_traits - hauteur:.1f} pt sous la page"
        assert bas_texte <= hauteur, f"page {page.number + 1} : texte à {bas_texte - hauteur:.1f} pt sous la page"


def test_cartouche_complet_dans_le_pdf(pdf_du_gabarit):
    _, _, pdf = pdf_du_gabarit
    for page in _pages_a_cartouche(pdf):
        texte = " ".join(page.get_text().split())
        for attendu in ("Ce plan est la propriété", "poursuites judiciaires",
                        "260 rue des Barronnières", "01700 - BEYNOST", "Tel. +33 4 78 88 14 00"):
            assert attendu in texte, f"page {page.number + 1} : « {attendu} » absent du PDF"


def test_adresse_sous_le_logo(pdf_du_gabarit):
    """Sans ses paragraphes vides de tête, l'adresse remontait sur le logo."""
    _, _, pdf = pdf_du_gabarit
    for page in _pages_a_cartouche(pdf):
        adresse = page.search_for("260 rue des")[0]
        # Tracé vectoriel dans LD82040, image dans LD64397.
        zones = [d["rect"] for d in page.get_drawings() if d.get("fill")] + \
                [fitz.Rect(i["bbox"]) for i in page.get_image_info()]
        logos = [r for r in zones if r.x1 < 200 and 515 < r.y0 and r.y1 < 550]
        assert logos, "logo du cartouche introuvable : le test ne prouverait rien"
        assert adresse.y0 >= max(r.y1 for r in logos)


def test_regle_des_cellules_fusionnees(pdf_du_gabarit):
    """Dans la copie : pour toute cellule fusionnée sur plusieurs lignes,
    paragraphes (vides compris) x interligne + marges <= hauteur nominale de
    la dernière ligne couverte."""
    _, copie, _ = pdf_du_gabarit
    with zipfile.ZipFile(copie) as z:
        slides = [n for n in z.namelist() if util._MOTIF_SLIDE.fullmatch(n)]
        controlees = 0
        for nom in slides:
            for tbl in etree.fromstring(z.read(nom)).iter(A + "tbl"):
                lignes = tbl.findall(A + "tr")
                for i, tr in enumerate(lignes):
                    for tc in tr.findall(A + "tc"):
                        portee = int(tc.get("rowSpan", "1"))
                        if portee <= 1:
                            continue
                        tcPr = tc.find(A + "tcPr")
                        marges = (int(tcPr.get("marT", 45720)) + int(tcPr.get("marB", 45720))) / cl.EMU_PAR_PT
                        contenu = sum(cl._hauteur_paragraphe_pt(p) for p in tc.iter(A + "p")) + marges
                        derniere = int(lignes[i + portee - 1].get("h")) / cl.EMU_PAR_PT
                        assert contenu <= derniere, f"{nom} : {contenu:.1f} pt > {derniere:.1f} pt"
                        controlees += 1
    assert controlees, "aucune cellule fusionnée : le test ne prouverait rien"


def test_la_copie_ne_deplace_ni_la_zone_du_dessin_ni_les_formes(pdf_du_gabarit, tmp_path):
    """Le haut du cartouche (d'où dérive la zone du dessin, cf.
    core/assemble.py::bornes_verticales) et la position de toute autre forme
    sont identiques dans la copie : seul le contenu des cellules change."""
    from pptx import Presentation

    gabarit, copie, _ = pdf_du_gabarit
    origine, adaptee = Presentation(str(gabarit)), Presentation(str(copie))
    for s_o, s_a in zip(list(origine.slides)[1:], list(adaptee.slides)[1:]):
        assert bornes_verticales(s_o, origine.slide_width) == bornes_verticales(s_a, adaptee.slide_width)
        assert [(sh.name, sh.left, sh.top, sh.width) for sh in s_o.shapes] == \
               [(sh.name, sh.left, sh.top, sh.width) for sh in s_a.shapes]


def test_slide_sans_tableau_inchangee():
    xml = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           b'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
    assert cl.ajuster_slide_xml(xml, 7556500) == xml
