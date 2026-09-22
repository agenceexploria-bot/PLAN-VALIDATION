# -*- coding: utf-8 -*-
"""
B3 — le plan a la même zone, donc la même position et une taille cohérente, sur
TOUTES les planches (slides 3 à N), assez grand pour être lisible.

Contrat de la zone standard, dérivé de la géométrie des gabarits (pas du code) :
  - sous le bandeau (dont le bas est à 2,10 cm) avec 0,2 cm de marge,
  - au-dessus du cartouche (dont le haut est à 17,97 cm) avec 0,2 cm de marge,
  - 0,6 cm de marge à gauche et à droite.
Toute vue extraite y tient en conservant son ratio (jamais étirée), centrée, et
touche la zone sur au moins un axe (elle ne reste pas « petite au milieu du blanc »).
Toute image d'origine du gabarit située dans cette bande est retirée — repérée
par sa POSITION, jamais par sa largeur.
"""
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Cm

from core import assemble

TYPES = ["non accompagné", "accompagné"]

ZONE_GAUCHE, ZONE_HAUT = Cm(0.6), Cm(2.10) + Cm(0.2)
ZONE_BAS = Cm(17.97) - Cm(0.2)
LARGEUR_SLIDE = Cm(29.70)
ZONE_LARGEUR = LARGEUR_SLIDE - 2 * Cm(0.6)
ZONE_HAUTEUR = ZONE_BAS - ZONE_HAUT
TOL = Cm(0.02)


def _images(tmp_path, ratios):
    """Une image par ratio largeur/hauteur (ex. A3 paysage 1.414, portrait 0.7…)."""
    planches = []
    for i, ratio in enumerate(ratios, 1):
        h = 600
        Image.new("RGB", (int(h * ratio), h), "white").save(tmp_path / f"p{i}.png")
        planches.append({"image": f"p{i}.png", "page_n": i, "labels": [], "page_w_pt": 1190.55})
    return planches


def _assembler(tmp_path, type_equipement, ratios):
    out = tmp_path / "zone.pptx"
    assemble.assembler({
        "base": assemble.base_pour_type(type_equipement), "out": out, "workdir": tmp_path,
        "meta": {"numero": "LDTEST700", "client": "C", "dessinateur": "AB", "indice": "R00",
                 "date": "01/01/2026", "equipement": "MONTE-CHARGE"},
        "specs": {}, "planches": _images(tmp_path, ratios),
    })
    return Presentation(out)


def _pictures_de_la_bande(slide):
    return [sh for sh in slide.shapes
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.top >= Cm(2.0) and sh.top + sh.height <= Cm(18.2)]


RATIOS = [1.414, 2.5, 1.0, 0.7]  # A3 paysage, très large, carré, portrait


@pytest.mark.parametrize("type_equipement", TYPES)
def test_une_seule_image_par_planche_dans_la_bande_de_travail(tmp_path, type_equipement):
    """Aucune image d'origine du gabarit ne subsiste (LD64397 slide 4 en portait
    une seconde, plus petite, qui échappait à l'ancien seuil `width > 8 cm`)."""
    prs = _assembler(tmp_path, type_equipement, RATIOS)
    for i in range(2, 2 + len(RATIOS)):
        assert len(_pictures_de_la_bande(prs.slides[i])) == 1, f"slide {i + 1} : image(s) d'origine restante(s)"


@pytest.mark.parametrize("type_equipement", TYPES)
def test_la_vue_reste_dans_la_zone_et_ne_deborde_ni_bandeau_ni_cartouche(tmp_path, type_equipement):
    prs = _assembler(tmp_path, type_equipement, RATIOS)
    for i in range(2, 2 + len(RATIOS)):
        (pic,) = _pictures_de_la_bande(prs.slides[i])
        assert pic.left >= ZONE_GAUCHE - TOL and pic.left + pic.width <= ZONE_GAUCHE + ZONE_LARGEUR + TOL
        assert pic.top >= ZONE_HAUT - TOL and pic.top + pic.height <= ZONE_BAS + TOL, f"slide {i + 1}"


@pytest.mark.parametrize("type_equipement", TYPES)
def test_la_vue_remplit_la_zone_sur_au_moins_un_axe_et_est_centree(tmp_path, type_equipement):
    prs = _assembler(tmp_path, type_equipement, RATIOS)
    for i, ratio in zip(range(2, 2 + len(RATIOS)), RATIOS):
        (pic,) = _pictures_de_la_bande(prs.slides[i])
        remplit_largeur = abs(pic.width - ZONE_LARGEUR) <= TOL
        remplit_hauteur = abs(pic.height - ZONE_HAUTEUR) <= TOL
        assert remplit_largeur or remplit_hauteur, f"slide {i + 1} : vue trop petite dans la zone"
        assert pic.width / pic.height == pytest.approx(ratio, rel=0.01), "ratio déformé"
        centre_x = pic.left + pic.width / 2
        centre_y = pic.top + pic.height / 2
        assert centre_x == pytest.approx(ZONE_GAUCHE + ZONE_LARGEUR / 2, abs=TOL)
        assert centre_y == pytest.approx(ZONE_HAUT + ZONE_HAUTEUR / 2, abs=TOL)


@pytest.mark.parametrize("type_equipement", TYPES)
def test_meme_ratio_meme_taille_meme_position_sur_toutes_les_planches(tmp_path, type_equipement):
    prs = _assembler(tmp_path, type_equipement, [1.414] * 4)
    geoms = []
    for i in range(2, 6):
        (pic,) = _pictures_de_la_bande(prs.slides[i])
        geoms.append((pic.left, pic.top, pic.width, pic.height))
    for g in geoms[1:]:
        assert all(abs(a - b) <= 2 for a, b in zip(g, geoms[0])), geoms


# ---------------------------------------------------------------------------
# Intégrité : une planche ne porte QUE ses propres étiquettes. Les planches au-delà
# des slides du gabarit sont des clones : ils étaient pris sur la slide 3 déjà
# décorée, et héritaient donc du plan ET des étiquettes de la planche 1.
# ---------------------------------------------------------------------------

def _planches_avec_etiquettes(tmp_path, n):
    """La planche i porte i étiquettes (donc chaque slide a un nombre distinct)."""
    planches = _images(tmp_path, [1.414] * n)
    for i, pl in enumerate(planches, 1):
        pl["labels"] = [{"text": f"P{i}-{k}", "bbox": [100 + 40 * k, 100, 130 + 40 * k, 115],
                         "vertical": False, "source": "X", "statut": "glossaire"} for k in range(i)]
    return planches


@pytest.mark.parametrize("type_equipement,n", [("non accompagné", 5), ("accompagné", 9)])
def test_chaque_planche_ne_porte_que_ses_propres_etiquettes(tmp_path, type_equipement, n):
    out = tmp_path / "clones.pptx"
    assemble.assembler({
        "base": assemble.base_pour_type(type_equipement), "out": out, "workdir": tmp_path,
        "meta": {"numero": "LDTEST800", "client": "C", "dessinateur": "AB", "indice": "R00",
                 "date": "01/01/2026", "equipement": "MONTE-CHARGE"},
        "specs": {}, "planches": _planches_avec_etiquettes(tmp_path, n),
    })
    import re
    prs = Presentation(out)
    assert len(prs.slides) == 2 + n
    for i in range(1, n + 1):
        # TOUTES les étiquettes de la slide (avant ou après l'image dans l'ordre d'empilement)
        textes = [sh.text_frame.text for sh in prs.slides[1 + i].shapes
                  if sh.has_text_frame and re.fullmatch(r"P\d+-\d+", sh.text_frame.text)]
        assert sorted(textes) == sorted(f"P{i}-{k}" for k in range(i)), f"planche {i} : {textes}"


# ---------------------------------------------------------------------------
# Le tableau specs de la slide 2 ne doit JAMAIS recouvrir le cartouche (audit §4 :
# hauteur de ligne fixe, sans limite). Constaté au rendu avec un vrai tableau de
# 20 lignes : il masquait « Date de création » et « Client ».
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("type_equipement", TYPES)
@pytest.mark.parametrize("n_lignes", [3, 20, 40])
def test_tableau_specs_reste_au_dessus_du_cartouche(tmp_path, type_equipement, n_lignes):
    out = tmp_path / "specs.pptx"
    assemble.assembler({
        "base": assemble.base_pour_type(type_equipement), "out": out, "workdir": tmp_path,
        "meta": {"numero": "LDTEST810", "client": "C", "dessinateur": "AB", "indice": "R00",
                 "date": "01/01/2026", "equipement": "MONTE-CHARGE"},
        "specs": {"table": [(f"Libellé {i}", f"Valeur {i}") for i in range(n_lignes)]},
        "planches": [],
    })
    slide = Presentation(out).slides[1]
    cartouche_top = min(sh.top for sh in slide.shapes if sh.has_table and len(sh.table.rows) == 4)
    (tableau,) = [sh for sh in slide.shapes if sh.has_table and len(sh.table.rows) == n_lignes]
    hauteur_lignes = sum(r.height for r in tableau.table.rows)
    assert tableau.top + hauteur_lignes <= cartouche_top, (
        f"{n_lignes} lignes : le tableau descend à {(tableau.top + hauteur_lignes) / 360000:.2f} cm, "
        f"le cartouche commence à {cartouche_top / 360000:.2f} cm"
    )
